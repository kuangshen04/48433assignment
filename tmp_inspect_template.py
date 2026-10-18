from openpyxl import load_workbook
import sys, textwrap

p = r"C:\Users\KuangShen\Desktop\uts\Software Architecture\Case Study Analysis Template.xlsx"
wb = load_workbook(p, data_only=False)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("\n" + "=" * 100)
    print(f"SHEET: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print("merged:", len(ws.merged_cells.ranges))
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200)):
        cells = []
        for c in row:
            if c.value is None:
                continue
            v = str(c.value).replace("\n", " / ")
            if len(v) > 300:
                v = v[:300] + "…"
            cells.append(f"{c.coordinate}={v}")
        if cells:
            print(" | ".join(cells))
