import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=False)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print("=" * 78)
    print(f"SHEET {ws.title!r}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print("merged:", [str(r) for r in ws.merged_cells.ranges][:20])
    for row in ws.iter_rows():
        vals = []
        for c in row:
            if c.value is not None and str(c.value).strip():
                vals.append(f"{c.coordinate}={str(c.value).strip()[:160]}")
        if vals:
            print(" | ".join(vals))
